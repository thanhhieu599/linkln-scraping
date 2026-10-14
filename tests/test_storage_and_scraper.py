from pathlib import Path

import pytest
from openpyxl import load_workbook

from linkedin_scraper import storage
from linkedin_scraper.config import Settings
from linkedin_scraper.models import Experience, Profile
from linkedin_scraper.scraper import LinkedInScraper, read_urls

FIXTURES = Path(__file__).parent / "fixtures"


def _settings(tmp_path: Path) -> Settings:
    s = Settings()
    s.state_path = tmp_path / "auth" / "state.json"
    s.output_jsonl = tmp_path / "profiles.jsonl"
    s.html_dir = tmp_path / "html"
    s.min_delay = s.max_delay = 0
    s.headless = True
    return s


def test_read_urls_dedupes_and_skips_comments(tmp_path):
    f = tmp_path / "urls.txt"
    f.write_text("# comment\nlinkedin.com/in/a\nhttps://www.linkedin.com/in/a/?x=1\n\nhttps://www.linkedin.com/in/b/\n")
    assert read_urls(f) == ["https://www.linkedin.com/in/a/", "https://www.linkedin.com/in/b/"]


def test_jsonl_resume_and_exports(tmp_path):
    path = tmp_path / "p.jsonl"
    storage.append_profile(path, Profile(url="u1", error="timeout"))
    storage.append_profile(path, Profile(url="u1", full_name="A",
                                         experiences=[Experience(title="Buyer", company="X")]))
    storage.append_profile(path, Profile(url="u2", error="boom"))
    assert storage.done_urls(path) == {"u1"}

    profiles = storage.load_profiles(path)
    assert len(profiles) == 2

    for ext in ("csv", "json", "xlsx"):
        assert storage.export(profiles, tmp_path / f"out.{ext}").exists()
    wb = load_workbook(tmp_path / "out.xlsx")
    assert wb.sheetnames == ["Profiles", "Experiences", "Educations"]
    assert wb["Experiences"]["C2"].value == "Buyer"

    with pytest.raises(ValueError):
        storage.export(profiles, tmp_path / "out.txt")


def test_end_to_end_with_mocked_linkedin(tmp_path, monkeypatch):
    """Chạy Playwright thật nhưng chặn mọi request tới linkedin.com và trả về HTML mẫu."""
    pytest.importorskip("playwright")
    from linkedin_scraper import scraper as scraper_mod

    settings = _settings(tmp_path)
    settings.state_path.parent.mkdir(parents=True)
    settings.state_path.write_text('{"cookies": [], "origins": []}')

    profile_html = (FIXTURES / "profile.html").read_text(encoding="utf-8")
    skills_html = (FIXTURES / "details_skills.html").read_text(encoding="utf-8")
    empty_details = "<html><body><main><ul></ul></main></body></html>"

    def handler(route):
        url = route.request.url
        if "/details/skills/" in url:
            body = skills_html
        elif "/details/" in url:
            body = empty_details
        elif "/in/" in url:
            body = profile_html
        else:
            return route.abort()
        route.fulfill(status=200, content_type="text/html; charset=utf-8", body=body)

    real_open = scraper_mod.open_context

    from contextlib import contextmanager

    @contextmanager
    def mocked_open(s, **kw):
        try:
            with real_open(s, **kw) as ctx:
                ctx.route("**/*", handler)
                yield ctx
        except Exception as exc:  # Chromium không khởi động được trên máy CI
            if "Executable doesn't exist" in str(exc):
                pytest.skip("Chromium chưa được cài")
            raise

    monkeypatch.setattr(scraper_mod, "open_context", mocked_open)
    monkeypatch.setattr(LinkedInScraper, "_sleep", lambda *a, **k: None)

    results = LinkedInScraper(settings, fetch_details=True).run(
        ["https://www.linkedin.com/in/nguyen-van-a/"]
    )
    assert len(results) == 1
    p = results[0]
    assert p.error == ""
    assert p.full_name == "Nguyễn Văn A"
    assert len(p.experiences) == 3  # trang chi tiết rỗng không ghi đè dữ liệu trang chính
    assert len(p.skills) == 5       # trang chi tiết kỹ năng có nhiều hơn -> được dùng

    # Lần chạy thứ 2: hồ sơ đã có -> bỏ qua
    again = LinkedInScraper(settings).run(["https://www.linkedin.com/in/nguyen-van-a/"])
    assert again == []
