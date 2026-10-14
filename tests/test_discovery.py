from contextlib import contextmanager
from pathlib import Path
from urllib.parse import parse_qs, urlparse

import pytest
from openpyxl import load_workbook

from linkedin_scraper import discovery as d
from linkedin_scraper import scraper as scraper_mod
from linkedin_scraper.cli import main
from linkedin_scraper.config import Settings
from linkedin_scraper.scraper import DiscoveryScraper

FIXTURES = Path(__file__).parent / "fixtures"


def _read(name: str) -> str:
    return (FIXTURES / name).read_text(encoding="utf-8")


# --------------------------------------------------------------------------- URLs
def test_people_search_url_has_vietnam_filter_and_page():
    url = d.people_search_url("an toàn lao động", page=3)
    q = parse_qs(urlparse(url).query)
    assert q["keywords"] == ["an toàn lao động"]
    assert q["geoUrn"] == ['["104195383"]']
    assert q["page"] == ["3"]
    assert "page" not in parse_qs(urlparse(d.people_search_url("x")).query)
    assert "geoUrn" not in parse_qs(urlparse(d.people_search_url("x", geo_urn="")).query)


# --------------------------------------------------------------------------- parsing
def test_parse_people_classic_layout_takes_everyone_with_profile():
    leads = {l.public_id: l for l in d.parse_people_search_html(_read("search_people_classic.html"))}
    assert set(leads) == {"nguyen-van-ehs", "tran-thi-sales"}  # "LinkedIn Member" không có link

    a = leads["nguyen-van-ehs"]
    assert a.full_name == "Nguyễn Văn EHS"
    assert a.linkedin_url == "https://www.linkedin.com/in/nguyen-van-ehs/"
    assert a.degree == "2nd"
    assert a.headline == "EHS Manager at Samsung Electronics Vietnam"
    assert a.location == "Bac Ninh, Vietnam"
    assert a.current == "EHS Manager at Samsung Electronics Vietnam"

    # Không lọc người bán: vẫn lấy
    b = leads["tran-thi-sales"]
    assert b.headline == "Sales Executive - Phân phối đồ bảo hộ lao động"
    assert b.location == "Ho Chi Minh City, Vietnam"
    assert b.degree == "3rd+"


def test_parse_people_new_layout_without_known_classes():
    leads = {l.public_id: l for l in d.parse_people_search_html(_read("search_people_new.html"))}
    assert set(leads) == {"le-minh-purchasing", "pham-safety"}
    lm = leads["le-minh-purchasing"]
    assert lm.full_name == "Lê Minh" and lm.degree == "1st"
    assert lm.headline == "Purchasing Manager | Jabil Vietnam"
    assert lm.location == "Ho Chi Minh City"
    ps = leads["pham-safety"]
    assert ps.headline == "Safety Officer tại Nhà máy thép Hòa Phát"
    assert ps.location == "Hai Duong, Vietnam"


def test_parse_posts_authors_dedup_and_skip_company_pages():
    leads = d.parse_posts_search_html(_read("search_posts.html"))
    assert [l.public_id for l in leads] == ["hoang-ehs-post"]
    l = leads[0]
    assert l.full_name == "Hoàng Văn Post"
    assert l.headline == "HSE Supervisor at Foxconn"
    assert len(l.post_snippets) == 2 and "EN 388" in l.post_snippets[0]


def test_detect_limit_and_empty():
    assert d.detect_search_limit(_read("search_limit.html"))
    assert not d.detect_search_limit(_read("search_people_classic.html"))
    assert d.has_no_results(_read("search_empty.html"))
    assert d.parse_people_search_html(_read("search_empty.html")) == []


# --------------------------------------------------------------------------- storage
def test_merge_save_load_export(tmp_path):
    store: dict = {}
    for lead in d.parse_people_search_html(_read("search_people_classic.html")):
        assert d.merge_lead(store, lead, "people:EHS")
    again = d.parse_people_search_html(_read("search_people_classic.html"))
    assert not any(d.merge_lead(store, lead, "people:HSE") for lead in again)
    assert store["nguyen-van-ehs"].sources == ["people:EHS", "people:HSE"]

    path = tmp_path / "leads.jsonl"
    d.save_leads(path, store)
    loaded = d.load_leads(path)
    assert loaded.keys() == store.keys()

    assert d.write_profiles_txt(tmp_path / "profiles.txt", loaded) == 2
    from linkedin_scraper.scraper import read_urls
    assert len(read_urls(tmp_path / "profiles.txt")) == 2

    ws = load_workbook(d.export_leads(loaded, tmp_path / "leads.xlsx"))["Leads"]
    assert ws.max_row == 3 and ws["A1"].value == "full_name"
    assert d.export_leads(loaded, tmp_path / "leads.csv").read_text(encoding="utf-8-sig").count("\n") == 3


# --------------------------------------------------------------------------- end-to-end
@pytest.fixture
def chromium_ok():
    from playwright.sync_api import sync_playwright

    s = Settings()
    try:
        with sync_playwright() as pw:
            pw.chromium.launch(executable_path=s.chromium_executable or None).close()
    except Exception as exc:
        pytest.skip(f"Chromium không khởi động được: {exc}")


def _setup(tmp_path, monkeypatch, pages_by_kw: dict[str, list[str]], calls: list[str]):
    settings = Settings()
    settings.state_path = tmp_path / "auth" / "state.json"
    settings.state_path.parent.mkdir(parents=True)
    settings.state_path.write_text('{"cookies": [], "origins": []}')
    settings.leads_jsonl = tmp_path / "leads.jsonl"
    settings.data_dir = tmp_path
    settings.min_delay = settings.max_delay = 0

    def handler(route):
        url = route.request.url
        calls.append(url)
        q = parse_qs(urlparse(url).query)
        if "/search/results/content/" in url:
            body = _read("search_posts.html")
        elif "/search/results/people/" in url:
            kw = q["keywords"][0]
            page = int(q.get("page", ["1"])[0])
            files = pages_by_kw.get(kw, [])
            body = _read(files[page - 1]) if page <= len(files) else _read("search_empty.html")
        else:
            return route.abort()
        route.fulfill(status=200, content_type="text/html; charset=utf-8", body=body)

    real_open = scraper_mod.open_context

    @contextmanager
    def mocked_open(s, **kw):
        with real_open(s, **kw) as ctx:
            ctx.route("**/*", handler)
            yield ctx

    monkeypatch.setattr(scraper_mod, "open_context", mocked_open)
    monkeypatch.setattr(DiscoveryScraper, "_sleep", lambda *a, **k: None)
    monkeypatch.setattr(DiscoveryScraper, "_scroll_to_bottom", lambda *a, **k: None)
    return settings


def test_discover_end_to_end(tmp_path, monkeypatch, chromium_ok):
    calls: list[str] = []
    settings = _setup(tmp_path, monkeypatch, {
        "EHS": ["search_people_classic.html", "search_people_new.html"],
        "Safety": ["search_people_new.html"],
    }, calls)
    store = DiscoveryScraper(settings).run_discover(["EHS", "Safety"], ["bảo hộ lao động"], pages=5)

    assert set(store) == {"nguyen-van-ehs", "tran-thi-sales", "le-minh-purchasing",
                          "pham-safety", "hoang-ehs-post"}
    # EHS: trang 1, 2, trang 3 rỗng -> dừng (không gọi trang 4, 5)
    ehs_pages = [u for u in calls if "keywords=EHS" in u]
    assert len(ehs_pages) == 3
    assert all("geoUrn" in u for u in calls if "/people/" in u)
    assert store["pham-safety"].sources == ["people:EHS", "people:Safety"]
    assert store["hoang-ehs-post"].sources == ["posts:bảo hộ lao động"]
    assert len(d.load_leads(settings.leads_jsonl)) == 5

    # Chạy lại: gộp, không nhân bản
    store2 = DiscoveryScraper(settings).run_discover(["EHS"], [], pages=1)
    assert len(store2) == 5


def test_discover_stops_on_search_limit_and_keeps_data(tmp_path, monkeypatch, chromium_ok):
    calls: list[str] = []
    settings = _setup(tmp_path, monkeypatch, {
        "A": ["search_people_classic.html", "search_limit.html", "search_people_new.html"],
        "B": ["search_people_new.html"],
    }, calls)
    store = DiscoveryScraper(settings).run_discover(["A", "B"], ["x"], pages=5)
    assert set(store) == {"nguyen-van-ehs", "tran-thi-sales"}
    assert not any("keywords=B" in u for u in calls)
    assert not any("/content/" in u for u in calls)
    assert len(d.load_leads(settings.leads_jsonl)) == 2


def test_discover_max_new(tmp_path, monkeypatch, chromium_ok):
    calls: list[str] = []
    settings = _setup(tmp_path, monkeypatch, {
        "A": ["search_people_classic.html", "search_people_new.html"],
    }, calls)
    store = DiscoveryScraper(settings).run_discover(["A"], ["x"], pages=5, max_new=2)
    assert len(store) == 2
    assert len([u for u in calls if "/people/" in u]) == 1


def test_cli_discover(tmp_path, monkeypatch, chromium_ok, capsys):
    from linkedin_scraper import cli

    calls: list[str] = []
    s = _setup(tmp_path, monkeypatch, {"EHS": ["search_people_classic.html"]}, calls)
    monkeypatch.setattr(cli, "Settings", lambda: s)  # không ghi vào data/ thật
    out = tmp_path / "leads.xlsx"
    assert main(["discover", "-k", "EHS", "--no-posts", "--pages", "2", "-o", str(out)]) == 0
    text = capsys.readouterr().out
    assert "Tổng 2 người" in text
    assert load_workbook(out)["Leads"].max_row == 3
    assert "nguyen-van-ehs" in (tmp_path / "profiles.txt").read_text(encoding="utf-8")


def test_cli_discover_unknown_preset():
    assert main(["discover", "--preset", "khong-co"]) == 2
