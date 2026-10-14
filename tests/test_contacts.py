from pathlib import Path

import pytest
from openpyxl import load_workbook

from linkedin_scraper import storage
from linkedin_scraper.cli import main
from linkedin_scraper.contacts import (
    Contact,
    best_email,
    clean_name,
    email_type,
    extract_emails,
    extract_phones,
    format_phone_local,
    is_valid_email,
    merge_contact_info,
    normalize_phone,
    parse_contact_info_html,
    read_connections_export,
    split_name,
)

FIXTURES = Path(__file__).parent / "fixtures"


def _read(name: str) -> str:
    return (FIXTURES / name).read_text(encoding="utf-8")


# --------------------------------------------------------------------------- email
@pytest.mark.parametrize("text,expected", [
    ("Liên hệ: a.nguyen@abc.com.vn để báo giá", ["a.nguyen@abc.com.vn"]),
    ("Email: john.doe+rfq@Intel.COM.", ["john.doe+rfq@intel.com"]),
    ("purchasing [at] jabil [dot] com", ["purchasing@jabil.com"]),
    ("sales(at)onsemi(dot)com", ["sales@onsemi.com"]),
    ("hr@samsung.com, hr@samsung.com; HR@SAMSUNG.COM", ["hr@samsung.com"]),
    ("logo@2x.png không phải email", []),
    ("đi họp at 9am tại nhà máy", []),
    ("a..b@x.com sai", []),
    ("", []),
])
def test_extract_emails(text, expected):
    assert extract_emails(text) == expected


def test_email_helpers():
    assert is_valid_email("a@b.co")
    assert not is_valid_email("a@b")
    assert not is_valid_email("@b.com")
    assert email_type("x@gmail.com") == "personal"
    assert email_type("x@yahoo.com.vn") == "personal"
    assert email_type("x@foxconn.com") == "business"
    assert best_email(["me@gmail.com", "me@foxconn.com"]) == "me@foxconn.com"
    assert best_email(["me@gmail.com"]) == "me@gmail.com"
    assert best_email([]) == ""


# --------------------------------------------------------------------------- phone
@pytest.mark.parametrize("raw,expected", [
    ("0912 345 678", "+84912345678"),
    ("0912.345.678", "+84912345678"),
    ("+84 912 345 678", "+84912345678"),
    ("84912345678", "+84912345678"),  # có mã nước 84 nhưng thiếu dấu +
    ("0084912345678", "+84912345678"),
    ("(+84) 903-111-222", "+84903111222"),
    ("024 3826 1234", "+842438261234"),   # máy bàn Hà Nội
    ("028 3829 5555", "+842838295555"),   # máy bàn TP.HCM
    ("+82 10-1234-5678", "+821012345678"),  # Hàn Quốc (chuyên gia Samsung)
    ("+1 (408) 765-8080", "+14087658080"),  # Mỹ (Intel)
    ("12345", ""),
    ("abc", ""),
    ("", ""),
])
def test_normalize_phone(raw, expected):
    assert normalize_phone(raw) == expected


def test_extract_phones_ignores_dates_and_years():
    text = ("Làm việc 2015 - 2018, ngày 12/05/2020. Hotline: 0912 345 678 "
            "hoặc Zalo 0987654321. Fax +84 24 3826 1234. Mã số thuế 0101234567")
    phones = extract_phones(text)
    assert "+84912345678" in phones
    assert "+84987654321" in phones
    assert "+842438261234" in phones
    assert all("2015" not in p and "2020" not in p for p in phones)


def test_format_phone_local():
    assert format_phone_local("+84912345678") == "0912 345 678"


# --------------------------------------------------------------------------- name
def test_names():
    assert split_name("Nguyễn Văn An") == ("An", "Nguyễn Văn")
    assert split_name("Madonna") == ("Madonna", "")
    assert split_name("") == ("", "")
    assert clean_name("Nguyễn Văn A, MBA") == "Nguyễn Văn A"
    assert clean_name("John Doe (He/Him)") == "John Doe"
    assert clean_name("Trần B | Purchasing") == "Trần B"


# --------------------------------------------------------------------------- contact-info overlay
def test_contact_info_classic_layout():
    info = parse_contact_info_html(_read("contact_info_classic.html"))
    assert info["emails"] == ["A.Nguyen@abc-electronics.vn"]
    assert info["phones"] == ["+84912345678", "+842438261234"]
    assert info["phone_types"] == ["Mobile", "Work"]
    assert info["websites"] == ["https://abc-electronics.vn"]
    assert info["twitter"] == "nguyenvana"
    assert info["address"] == "KCN Quế Võ, Bắc Ninh"
    assert info["connected_on"] == "March 5, 2021"


def test_contact_info_new_layout_without_mailto():
    info = parse_contact_info_html(_read("contact_info_new.html"))
    assert info["emails"] == ["tran.b@gmail.com"]
    assert info["phones"] == ["+84903111222"]
    assert info["phone_types"] == ["Mobile"]
    assert info["connected_on"] == "2 Jan 2023"


def test_contact_info_not_shared():
    info = parse_contact_info_html(_read("contact_info_empty.html"))
    assert info["emails"] == [] and info["phones"] == []


def test_merge_dedupes_and_tracks_source():
    c = Contact(linkedin_url="u", full_name="Nguyễn Văn A")
    merge_contact_info(c, {"emails": ["a@abc.vn"], "phones": ["+84912345678"]}, "about")
    merge_contact_info(c, {"emails": ["A@abc.vn", "b@abc.vn"], "phones": ["+84912345678"],
                           "phone_types": ["Mobile"]}, "contact_info")
    merge_contact_info(c, {"emails": ["a@abc.vn"]}, "noop")
    assert c.emails == ["a@abc.vn", "b@abc.vn"]
    assert c.phones == ["+84912345678"]
    assert c.sources == ["about", "contact_info"]
    row = c.to_row()
    assert row["first_name"] == "A" and row["last_name"] == "Nguyễn Văn"
    assert row["email"] == "a@abc.vn" and row["email_type"] == "business"


# --------------------------------------------------------------------------- LinkedIn data export
def test_read_connections_export():
    contacts = read_connections_export(FIXTURES / "Connections.csv")
    assert len(contacts) == 4
    an = contacts[0]
    assert an.full_name == "Nguyễn Văn An"
    assert an.emails == ["an.nv@samsung.com"]
    assert an.company == "Samsung Electronics Vietnam"
    assert an.title == "Purchasing Manager"
    assert an.linkedin_url == "https://www.linkedin.com/in/nguyen-van-an/"
    assert an.sources == ["connections_export"]
    assert contacts[1].full_name == "John Smith" and contacts[1].emails == []
    assert contacts[2].company == "Jabil Vietnam, Inc."
    assert contacts[3].emails == []  # email sai định dạng bị loại


def test_read_connections_export_rejects_wrong_file(tmp_path):
    f = tmp_path / "x.csv"
    f.write_text("a,b\n1,2\n")
    with pytest.raises(ValueError):
        read_connections_export(f)


# --------------------------------------------------------------------------- export + CLI
def test_export_contacts_xlsx_keeps_phone_as_text(tmp_path):
    contacts = [
        Contact(linkedin_url="u1", full_name="Nguyễn Văn A", emails=["a@abc.vn"],
                phones=["+84912345678"], phone_types=["Mobile"], company="ABC"),
        Contact(linkedin_url="u2", full_name="Trần B"),
    ]
    out = storage.export_contacts(contacts, tmp_path / "c.xlsx")
    ws = load_workbook(out)["Contacts"]
    headers = [c.value for c in ws[1]]
    row = dict(zip(headers, [c.value for c in ws[2]]))
    assert row["email"] == "a@abc.vn"
    assert row["phone"] == "+84912345678"
    assert row["phone_local"] == "0912 345 678"
    assert row["first_name"] == "A"
    assert ws.max_row == 3

    csv_text = storage.export_contacts(contacts, tmp_path / "c.csv").read_text(encoding="utf-8-sig")
    assert csv_text.splitlines()[0].startswith("full_name,first_name,last_name,email")
    assert "+84912345678" in csv_text


def test_cli_import_connections_and_export(tmp_path, monkeypatch, capsys):
    from linkedin_scraper import config

    monkeypatch.setattr(config.Settings, "contacts_jsonl", tmp_path / "contacts.jsonl")
    out = tmp_path / "out.xlsx"
    assert main(["import-connections", str(FIXTURES / "Connections.csv"), "--merge",
                 "-o", str(out), "--only-with-contact"]) == 0
    assert "4 kết nối, 2 có email" in capsys.readouterr().out
    assert load_workbook(out)["Contacts"].max_row == 3  # header + 2 người có email

    # Import lại không tạo trùng
    assert main(["import-connections", str(FIXTURES / "Connections.csv"), "--merge"]) == 0
    assert "Đã thêm 0 kết nối mới" in capsys.readouterr().out

    out2 = tmp_path / "all.csv"
    assert main(["export-contacts", "-o", str(out2)]) == 0
    assert len(out2.read_text(encoding="utf-8-sig").strip().splitlines()) == 5


def test_cli_missing_input_file_gives_friendly_message(tmp_path):
    with pytest.raises(SystemExit) as exc:
        main(["scrape", "-i", str(tmp_path / "khong-co.txt")])
    assert "Không tìm thấy file danh sách" in str(exc.value)
    with pytest.raises(SystemExit) as exc:
        main(["contacts", "-i", str(tmp_path / "khong-co.txt")])
    assert "profiles.example.txt" in str(exc.value)
