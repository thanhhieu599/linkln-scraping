"""Lưu trữ kết quả (JSONL, chạy lại không trùng) và xuất CSV / Excel / JSON."""
from __future__ import annotations

import csv
import json
from pathlib import Path

from .models import Profile


def load_profiles(path: Path) -> list[Profile]:
    if not path.exists():
        return []
    latest: dict[str, Profile] = {}
    with path.open(encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            profile = Profile.from_dict(json.loads(line))
            latest[profile.url] = profile  # bản ghi sau ghi đè bản trước
    return list(latest.values())


def done_urls(path: Path) -> set[str]:
    return {p.url for p in load_profiles(path) if not p.error}


def append_profile(path: Path, profile: Profile) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as f:
        f.write(json.dumps(profile.to_dict(), ensure_ascii=False) + "\n")


def export(profiles: list[Profile], out: Path) -> Path:
    out.parent.mkdir(parents=True, exist_ok=True)
    suffix = out.suffix.lower()
    if suffix == ".json":
        out.write_text(
            json.dumps([p.to_dict() for p in profiles], ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
    elif suffix == ".csv":
        rows = [p.to_flat_row() for p in profiles]
        # utf-8-sig để Excel mở tiếng Việt không lỗi font
        with out.open("w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=list(Profile(url="").to_flat_row().keys()))
            writer.writeheader()
            writer.writerows(rows)
    elif suffix == ".xlsx":
        _export_xlsx(profiles, out)
    else:
        raise ValueError(f"Định dạng không hỗ trợ: {suffix} (dùng .csv, .xlsx hoặc .json)")
    return out


def _export_xlsx(profiles: list[Profile], out: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Profiles"
    headers = list(Profile(url="").to_flat_row().keys())
    ws.append(headers)
    for p in profiles:
        ws.append([p.to_flat_row()[h] for h in headers])

    ws_exp = wb.create_sheet("Experiences")
    exp_headers = ["url", "full_name", "title", "company", "employment_type",
                   "date_range", "duration", "location", "description"]
    ws_exp.append(exp_headers)
    for p in profiles:
        for e in p.experiences:
            ws_exp.append([p.url, p.full_name, e.title, e.company, e.employment_type,
                           e.date_range, e.duration, e.location, e.description])

    ws_edu = wb.create_sheet("Educations")
    ws_edu.append(["url", "full_name", "school", "degree", "date_range", "description"])
    for p in profiles:
        for e in p.educations:
            ws_edu.append([p.url, p.full_name, e.school, e.degree, e.date_range, e.description])

    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
    wb.save(out)
